import base64
import hashlib
import html
import io
import json
import os
import uuid
from datetime import datetime

import streamlit as st
from langchain_core.messages import AIMessage, HumanMessage, SystemMessage
from langchain_openai import ChatOpenAI

APP_NAME = "Custom ChatGPT"
DEFAULT_SYSTEM = "You are a helpful assistant."
DEFAULT_MODEL = "gpt-5.2"
USER_AVATAR = ":material/person:"
BOT_AVATAR = ":material/auto_awesome:"

st.set_page_config(page_title=APP_NAME, page_icon=":material/auto_awesome:", layout="wide")

SUGGESTIONS = {
    "Popular": [
        ("tips_and_updates", "Give me 5 ways to study more effectively",
         "Give me 5 practical, science-backed ways to study more effectively, with a short example for each."),
        ("psychology", "Explain how AI chatbots like ChatGPT work",
         "Explain how AI chatbots like ChatGPT work, in simple terms, using an analogy."),
        ("event_note", "Plan a balanced weekly routine for me",
         "Help me plan a balanced weekly routine for study, exercise and rest. Give it as a table."),
        ("mail", "Write a polite email asking for an extension",
         "Write a short, polite email to a teacher asking for a deadline extension."),
    ],
    "Learn": [
        ("school", "Explain recursion with a simple example",
         "Explain recursion to a beginner with a simple example and a diagram made of text."),
        ("public", "How does the internet actually work?",
         "How does the internet actually work? Explain step by step, from typing a URL to seeing a page."),
        ("hub", "AI vs machine learning vs deep learning",
         "What is the difference between AI, machine learning and deep learning? Use a simple comparison table."),
        ("database", "Teach me SQL joins with examples",
         "Teach me the different SQL joins with small example tables and results."),
    ],
    "Code": [
        ("terminal", "Build a small to-do app in Python",
         "Show me how to build a small command-line to-do app in Python, with clean code and comments."),
        ("functions", "Explain Big O notation with examples",
         "Explain Big O notation with short code examples for O(1), O(n) and O(n^2)."),
        ("table_view", "Read and summarize a CSV file with pandas",
         "Show me how to read a CSV file with pandas and print a quick summary of it."),
        ("bug_report", "Common beginner mistakes in Python",
         "What are the most common beginner mistakes in Python, and how do I avoid each one?"),
    ],
    "Write": [
        ("badge", "Write a LinkedIn summary for a fresh graduate",
         "Write a professional LinkedIn summary for a fresh IT graduate looking for their first job."),
        ("description", "Write a cover letter for a junior developer role",
         "Write a concise cover letter for a junior developer role. Leave placeholders for my details."),
        ("edit_note", "Write a thank-you note after an interview",
         "Write a short, warm thank-you email to send after a job interview."),
        ("smart_display", "10 catchy YouTube Shorts title ideas",
         "Give me 10 catchy YouTube Shorts title ideas about learning to code."),
    ],
    "Career": [
        ("record_voice_over", "Practice a junior developer interview",
         "Act as an interviewer. Ask me 5 common junior developer interview questions, one at a time, and give feedback on each answer."),
        ("security", "How do I become a cyber security analyst?",
         "What skills and certifications should I learn to become a cyber security analyst? Give me a roadmap."),
        ("work", "Build a strong CV with no experience",
         "How do I build a strong CV when I have no work experience? Give me a section-by-section guide."),
        ("payments", "How to negotiate a first salary",
         "How do I negotiate my first salary politely and confidently? Include example sentences."),
    ],
    "Fun": [
        ("lightbulb", "Tell me a fun fact I probably don't know",
         "Tell me a surprising fun fact I probably don't know, and explain why it's true."),
        ("extension", "Give me a riddle to solve",
         "Give me a clever riddle to solve. Don't reveal the answer until I ask."),
        ("celebration", "Suggest a weekend project to build my skills",
         "Suggest 3 fun weekend projects I can build to improve my programming skills."),
        ("menu_book", "Recommend 5 books for curious people",
         "Recommend 5 books for curious people, with one line on why each is worth reading."),
    ],
}

THEMES = {
    "dark": {
        "bg": "#0e0a1c",
        "panel": "rgba(20,15,38,0.82)",
        "surface": "#171230",
        "card": "rgba(255,255,255,0.05)",
        "card-hover": "rgba(255,255,255,0.10)",
        "line": "rgba(255,255,255,0.10)",
        "text": "#f5f2ff",
        "muted": "#a49dc6",
        "faint": "#6f6893",
        "accent": "#9b7bff",
        "orb1": "rgba(124,77,255,0.50)",
        "orb2": "rgba(255,138,101,0.24)",
        "orb3": "rgba(56,189,248,0.20)",
        "shadow": "0 18px 50px rgba(0,0,0,0.45)",
        "scheme": "dark",
    },
    "light": {
        "bg": "#f7f4ff",
        "panel": "rgba(255,255,255,0.80)",
        "surface": "#ffffff",
        "card": "rgba(255,255,255,0.72)",
        "card-hover": "#ffffff",
        "line": "#e6defa",
        "text": "#1c1730",
        "muted": "#716a8e",
        "faint": "#a29bbe",
        "accent": "#7c4dff",
        "orb1": "rgba(124,77,255,0.28)",
        "orb2": "rgba(255,154,118,0.36)",
        "orb3": "rgba(120,200,255,0.26)",
        "shadow": "0 18px 50px rgba(124,77,255,0.16)",
        "scheme": "light",
    },
}

BASE_CSS = """
@import url('https://fonts.googleapis.com/css2?family=Bricolage+Grotesque:opsz,wght@12..96,500..800&family=DM+Sans:wght@400;500;600&family=JetBrains+Mono:wght@400;500&display=swap');

:root { --g1: #7c4dff; --g2: #d946ef; --g3: #ff9a76; --g4: #38bdf8; }

@property --a { syntax: '<angle>'; inherits: false; initial-value: 0deg; }

/* ================= base ================= */
html, body, .stApp, [data-testid="stAppViewContainer"], button, input, textarea, select {
    font-family: 'DM Sans', system-ui, sans-serif !important;
}
pre, code, [data-testid="stCode"] * { font-family: 'JetBrains Mono', ui-monospace, monospace !important; }
[data-testid="stIconMaterial"], .material-symbols-rounded { font-family: 'Material Symbols Rounded' !important; }

[data-testid="stAppDeployButton"], [data-testid="stMainMenu"],
[data-testid="stDecoration"], footer { display: none !important; }
[data-testid="stAppViewContainer"], [data-testid="stHeader"] {
    position: relative; z-index: 1; background: transparent !important;
}
[data-testid="stExpandSidebarButton"], [data-testid="stSidebarCollapseButton"],
[data-testid="collapsedControl"], [data-testid="stSidebarCollapsedControl"] {
    display: flex !important; visibility: visible !important; opacity: 1 !important;
}
[data-testid="stExpandSidebarButton"] *, [data-testid="stSidebarCollapseButton"] *,
[data-testid="collapsedControl"] *, [data-testid="stSidebarCollapsedControl"] * {
    color: var(--accent) !important;
}

.stApp { background: var(--bg) !important; color: var(--text); }

.stApp::before {
    content: ""; position: fixed; inset: -20%; z-index: 0; pointer-events: none;
    background:
        radial-gradient(38% 38% at 18% 22%, var(--orb1), transparent 70%),
        radial-gradient(34% 34% at 82% 18%, var(--orb2), transparent 70%),
        radial-gradient(40% 40% at 64% 86%, var(--orb3), transparent 70%);
    filter: blur(30px);
    animation: drift 30s ease-in-out infinite alternate;
}

::-webkit-scrollbar { width: 8px; height: 8px; }
::-webkit-scrollbar-thumb { background: var(--line); border-radius: 8px; }

.block-container { max-width: 800px; padding-top: 1.5rem; padding-bottom: 7rem; }
[data-testid="stBottom"], [data-testid="stBottom"] > div { background: transparent !important; }
[data-testid="stBottomBlockContainer"] { max-width: 800px; margin: 0 auto; }

.stApp label, .stApp [data-testid="stWidgetLabel"] p { color: var(--muted) !important; font-weight: 500; }
[data-testid="stCaptionContainer"], .stApp .stCaption { color: var(--faint) !important; }

/* ================= the orb ================= */
.orb {
    --s: 76px; position: relative; flex: none;
    width: var(--s); height: var(--s); border-radius: 50%;
    background: conic-gradient(from var(--a), var(--g1), var(--g2), var(--g3), var(--g4), var(--g1));
    box-shadow: 0 0 calc(var(--s) * 0.55) rgba(124,77,255,0.55);
    animation: spin 7s linear infinite, breathe 4.5s ease-in-out infinite;
}
.orb::after {
    content: ""; position: absolute; inset: 7%; border-radius: 50%;
    background:
        radial-gradient(circle at 32% 26%, rgba(255,255,255,0.9), rgba(255,255,255,0) 52%),
        radial-gradient(circle at 50% 50%, rgba(20,10,50,0.28), rgba(20,10,50,0) 70%);
}
.orb.sm { --s: 30px; }

/* ================= sidebar ================= */
[data-testid="stSidebar"] {
    background: var(--panel) !important;
    backdrop-filter: blur(18px) saturate(140%);
    border-right: 1px solid var(--line);
}
[data-testid="stSidebar"] > div { background: transparent !important; }
[data-testid="stSidebar"] [data-testid="stHorizontalBlock"] { gap: 0.2rem; align-items: center; }

.sb-brand { display: flex; align-items: center; gap: 0.7rem; margin: 0.2rem 0 1rem 0.2rem; }
.sb-name {
    font-family: 'Bricolage Grotesque', 'DM Sans', sans-serif; font-size: 1.2rem;
    font-weight: 700; letter-spacing: -0.02em; color: var(--text);
}
.sb-group { font-size: 0.76rem; font-weight: 600; color: var(--faint); margin: 1rem 0 0.25rem 0.6rem; }

[data-testid="stSidebar"] .stButton > button {
    width: 100%; justify-content: flex-start; text-align: left;
    background: transparent; border: 1px solid transparent; box-shadow: none;
    color: var(--muted); border-radius: 12px; padding: 0.42rem 0.7rem;
    font-size: 0.92rem; font-weight: 500;
    transition: background 0.2s ease, color 0.2s ease, transform 0.2s ease;
}
[data-testid="stSidebar"] .stButton > button p {
    overflow: hidden; text-overflow: ellipsis; white-space: nowrap; color: inherit !important;
}
[data-testid="stSidebar"] .stButton > button:hover {
    background: var(--card-hover); color: var(--text); transform: translateX(3px);
}
.st-key-active_chat_btn button {
    background: var(--card-hover) !important; color: var(--text) !important;
    border-color: var(--line) !important; font-weight: 600 !important;
    box-shadow: inset 3px 0 0 var(--accent), 0 0 18px rgba(124,77,255,0.18) !important;
}

.st-key-new_chat_btn button {
    position: relative; overflow: hidden; justify-content: center !important;
    background: linear-gradient(120deg, var(--g1), var(--g2)) !important;
    border: none !important; border-radius: 14px !important; padding: 0.62rem 1rem !important;
    box-shadow: 0 10px 24px rgba(124,77,255,0.35) !important;
}
.st-key-new_chat_btn button, .st-key-new_chat_btn button * {
    color: #fff !important; -webkit-text-fill-color: #fff !important; font-weight: 600;
}
.st-key-new_chat_btn button::after {
    content: ""; position: absolute; top: 0; left: -60%; width: 40%; height: 100%;
    background: linear-gradient(100deg, transparent, rgba(255,255,255,0.5), transparent);
    transform: skewX(-20deg);
}
.st-key-new_chat_btn button:hover { transform: translateY(-2px) !important; box-shadow: 0 14px 30px rgba(124,77,255,0.5) !important; }
.st-key-new_chat_btn button:hover::after { animation: shine 0.8s ease; }

/* chat row menu (the three dots) */
[data-testid="stSidebar"] [data-testid="stPopover"] button {
    background: transparent !important; border: none !important; box-shadow: none !important;
    color: var(--faint) !important; padding: 0.3rem 0 !important; min-height: 2rem;
    justify-content: center !important; opacity: 0.55; font-size: 1.15rem; transition: opacity 0.2s, color 0.2s;
}
[data-testid="stSidebar"] [data-testid="stPopover"] button:hover { opacity: 1; color: var(--text) !important; transform: none; }
[data-testid="stSidebar"] [data-testid="stPopover"] button [data-testid="stIconMaterial"] { display: none; }

/* ================= POPOVER FIX ================= */
[data-testid="stPopoverBody"] {
    background: var(--surface) !important; border: 1px solid var(--line) !important;
    border-radius: 16px !important; box-shadow: var(--shadow) !important; color: var(--text);
}
[data-testid="stPopoverBody"] .stButton > button {
    width: 100%; justify-content: flex-start; background: var(--card); border: 1px solid var(--line);
    color: var(--text); border-radius: 10px; transition: all 0.2s ease;
}
[data-testid="stPopoverBody"] .stButton > button:hover { border-color: var(--accent); color: var(--accent); }
[class*="st-key-del_"] button:hover { color: #ff6b81 !important; border-color: #ff6b81 !important; background: rgba(255,107,129,0.1) !important; }

/* Popover ke andar input ka text + placeholder + cursor dono visible */
[data-testid="stPopoverBody"] input,
[data-testid="stPopoverBody"] textarea,
[data-testid="stPopoverBody"] [data-baseweb="input"] input,
[data-testid="stPopoverBody"] [data-baseweb="base-input"] input {
    color: var(--text) !important;
    -webkit-text-fill-color: var(--text) !important;
    background: var(--surface) !important;
    opacity: 1 !important;
    caret-color: var(--accent) !important;
}
[data-testid="stPopoverBody"] input::placeholder,
[data-testid="stPopoverBody"] textarea::placeholder {
    color: var(--faint) !important;
    -webkit-text-fill-color: var(--faint) !important;
    opacity: 1 !important;
}

/* inputs */
[data-baseweb="input"], [data-baseweb="base-input"], [data-baseweb="textarea"], [data-baseweb="select"] > div {
    background: var(--card) !important; border-color: var(--line) !important; border-radius: 12px !important;
}
[data-baseweb="base-input"] { border: none !important; }
[data-baseweb="input"] { border: 1px solid var(--line) !important; outline: none !important; box-shadow: none !important; }
[data-baseweb="input"]:focus-within, [data-baseweb="textarea"]:focus-within {
    border-color: var(--accent) !important; box-shadow: 0 0 0 3px rgba(124,77,255,0.2) !important;
}
[data-testid="stTextInput"] input, [data-testid="stTextArea"] textarea {
    color: var(--text) !important; -webkit-text-fill-color: var(--text) !important; background: transparent !important;
}
input::placeholder, textarea::placeholder { color: var(--faint) !important; -webkit-text-fill-color: var(--faint) !important; }
[data-testid="stExpander"] { border: 1px solid var(--line) !important; border-radius: 14px !important; background: var(--card); }
[data-testid="stExpander"] details { background: transparent !important; }
[data-testid="stExpander"] summary, [data-testid="stExpander"] summary * { color: var(--text) !important; }

/* user chip + logout */
.user-chip {
    display: flex; align-items: center; gap: 0.65rem; padding: 0.55rem 0.65rem; margin-top: 0.7rem;
    border: 1px solid var(--line); border-radius: 14px; background: var(--card);
}
.user-chip .av {
    width: 32px; height: 32px; border-radius: 50%; display: flex; align-items: center; justify-content: center;
    background: linear-gradient(135deg, var(--g1), var(--g3)); color: #fff; font-weight: 700; font-size: 0.9rem;
}
.user-chip .nm { font-weight: 600; font-size: 0.9rem; color: var(--text); line-height: 1.2; }
.user-chip .sub { font-size: 0.75rem; color: var(--muted); }
.st-key-logout_btn button { color: var(--faint) !important; }
.st-key-logout_btn button:hover { color: #ff6b81 !important; }

/* ================= hero ================= */
.hero { display: flex; flex-direction: column; align-items: center; text-align: center; gap: 1.1rem; margin: 8vh 0 1.6rem 0; }
.hero .orb { --s: 84px; animation: orbIn 0.9s cubic-bezier(0.2, 0.9, 0.3, 1.2) both, spin 7s linear infinite, breathe 4.5s ease-in-out 1s infinite; }
.hero-title {
    font-family: 'Bricolage Grotesque', 'DM Sans', sans-serif; font-weight: 700;
    font-size: clamp(2rem, 4.8vw, 3.1rem); letter-spacing: -0.035em; line-height: 1.08;
    background: linear-gradient(100deg, var(--text) 0%, var(--text) 30%, var(--g1) 50%, var(--g2) 68%, var(--g3) 84%, var(--text) 100%);
    background-size: 260% 100%;
    -webkit-background-clip: text; background-clip: text;
    color: transparent !important; -webkit-text-fill-color: transparent;
    animation: riseIn 0.7s 0.1s cubic-bezier(0.2, 0.8, 0.2, 1) both, sheen 10s ease-in-out 1s infinite alternate;
}
.hero p { color: var(--muted); font-size: 1.05rem; margin: 0; animation: riseIn 0.7s 0.22s cubic-bezier(0.2, 0.8, 0.2, 1) both; }
.fine { text-align: center; color: var(--faint); font-size: 0.78rem; margin-top: 1.4rem; }

/* topic pills */
.st-key-sg_cat { display: flex; justify-content: center; margin: 1.3rem 0 0.4rem 0; }
.st-key-sg_cat [data-testid="stPills"] { justify-content: center; }
.st-key-sg_cat button {
    background: transparent !important; border: 1px solid var(--line) !important; color: var(--muted) !important;
    border-radius: 999px !important; font-weight: 500; transition: all 0.22s ease;
}
.st-key-sg_cat button:hover { border-color: var(--accent) !important; color: var(--text) !important; transform: translateY(-1px); }
.st-key-sg_cat button[data-testid="stBaseButton-pillsActive"] {
    background: linear-gradient(120deg, var(--g1), var(--g2)) !important; border-color: transparent !important;
    color: #fff !important; box-shadow: 0 8px 20px rgba(124,77,255,0.35);
}
.st-key-sg_cat button[data-testid="stBaseButton-pillsActive"] * { color: #fff !important; }

/* suggestion cards */
[class*="st-key-sg_card_"] button {
    width: 100%; min-height: 78px; justify-content: flex-start; align-items: flex-start; text-align: left;
    gap: 0.7rem; padding: 0.95rem 1rem; border-radius: 18px;
    background: var(--card); border: 1px solid var(--line); box-shadow: none;
    color: var(--text); font-weight: 500; line-height: 1.35; backdrop-filter: blur(8px);
    transition: transform 0.28s cubic-bezier(0.2, 0.8, 0.2, 1), background 0.2s, border-color 0.2s, box-shadow 0.28s;
    animation: riseIn 0.55s cubic-bezier(0.2, 0.8, 0.2, 1) both;
}
[class*="st-key-sg_card_"] button p { text-align: left; color: var(--text) !important; }
[class*="st-key-sg_card_0_"] button { animation-delay: 0.04s; }
[class*="st-key-sg_card_1_"] button { animation-delay: 0.11s; }
[class*="st-key-sg_card_2_"] button { animation-delay: 0.18s; }
[class*="st-key-sg_card_3_"] button { animation-delay: 0.25s; }
[class*="st-key-sg_card_"] button:hover {
    transform: translateY(-3px); background: var(--card-hover); border-color: var(--accent);
    box-shadow: 0 16px 34px rgba(124,77,255,0.22);
}
[class*="st-key-sg_card_"] button [data-testid="stIconMaterial"] {
    color: var(--accent); font-size: 1.4rem; transition: transform 0.35s cubic-bezier(0.2, 0.8, 0.2, 1.4);
}
[class*="st-key-sg_card_"] button:hover [data-testid="stIconMaterial"] { transform: scale(1.2) rotate(-8deg); }

/* ================= messages ================= */
[data-testid="stChatMessage"] { background: transparent !important; border: none; box-shadow: none; padding: 0.5rem 0; gap: 0.85rem; }
[data-testid="stChatMessage"] :is(p, li, h1, h2, h3, h4, strong, em, td, th, blockquote) { color: var(--text); }
[data-testid="stChatMessage"] :is(p, li) { line-height: 1.72; font-size: 0.99rem; }
[data-testid="stChatMessage"] :not(pre) > code {
    background: var(--card-hover); color: var(--accent) !important; border-radius: 6px; padding: 0.12rem 0.4rem; font-size: 0.88em;
}
[data-testid="stChatMessage"] pre { background: #14102a !important; border: 1px solid rgba(255,255,255,0.08); border-radius: 14px !important; }
[data-testid="stChatMessage"] table { border-collapse: collapse; }
[data-testid="stChatMessage"] :is(th, td) { border: 1px solid var(--line) !important; padding: 0.45rem 0.7rem; }

[data-testid="stChatMessageAvatarAssistant"] {
    background: conic-gradient(from var(--a), var(--g1), var(--g2), var(--g3), var(--g4), var(--g1)) !important;
    border-radius: 50% !important; box-shadow: 0 0 16px rgba(124,77,255,0.5); animation: spin 9s linear infinite;
}
[data-testid="stChatMessageAvatarAssistant"] * { color: #fff !important; }

[data-testid="stChatMessageAvatarUser"] { display: none !important; }
[data-testid="stChatMessage"]:has([data-testid="stChatMessageAvatarUser"]) {
    flex-direction: row-reverse; margin-left: auto; width: fit-content; max-width: 82%;
    background: linear-gradient(135deg, var(--g1), #a74cf0) !important;
    border-radius: 22px 22px 6px 22px; padding: 0.7rem 1.1rem;
    box-shadow: 0 10px 24px rgba(124,77,255,0.3);
}
[data-testid="stChatMessage"]:has([data-testid="stChatMessageAvatarUser"]) * { color: #fff !important; }

[class*="st-key-live_"] [data-testid="stChatMessage"] { animation: riseIn 0.5s cubic-bezier(0.2, 0.8, 0.2, 1) both; }
.st-key-live_user [data-testid="stChatMessage"] { animation-name: bubbleIn; transform-origin: right bottom; }

.thinking { display: flex; align-items: center; gap: 0.65rem; padding: 0.45rem 0; }
.thinking .dots { display: flex; gap: 5px; }
.thinking .dots i {
    width: 7px; height: 7px; border-radius: 50%; display: block;
    background: linear-gradient(135deg, var(--g1), var(--g3)); animation: wave 1.1s infinite ease-in-out;
}
.thinking .dots i:nth-child(2) { animation-delay: 0.14s; }
.thinking .dots i:nth-child(3) { animation-delay: 0.28s; }
.thinking .label {
    font-size: 0.92rem;
    background: linear-gradient(90deg, var(--muted) 40%, var(--text) 50%, var(--muted) 60%);
    background-size: 250% 100%; -webkit-background-clip: text; background-clip: text;
    color: transparent; -webkit-text-fill-color: transparent; animation: shimmer 1.8s linear infinite;
}

.st-key-regen_btn button {
    background: transparent !important; border: 1px solid var(--line) !important; color: var(--muted) !important;
    border-radius: 999px !important; padding: 0.25rem 0.85rem !important; font-size: 0.85rem; transition: all 0.2s ease;
}
.st-key-regen_btn button:hover { color: var(--accent) !important; border-color: var(--accent) !important; }
.st-key-regen_btn button [data-testid="stIconMaterial"] { transition: transform 0.5s ease; }
.st-key-regen_btn button:hover [data-testid="stIconMaterial"] { transform: rotate(360deg); }

/* ================= chat input ================= */
[data-testid="stChatInput"] {
    border-radius: 28px !important; background: var(--surface) !important;
    border: 1.5px solid var(--line) !important; box-shadow: var(--shadow);
    transition: box-shadow 0.3s ease, transform 0.3s ease;
}
[data-testid="stChatInput"] > div, [data-testid="stChatInput"] [data-baseweb="textarea"],
[data-testid="stChatInput"] [data-baseweb="base-input"] {
    background: transparent !important; border: none !important; box-shadow: none !important; outline: none !important;
}
[data-testid="stChatInput"] textarea {
    background: transparent !important; color: var(--text) !important;
    -webkit-text-fill-color: var(--text) !important; font-size: 1rem;
}
[data-testid="stChatInput"]:focus-within {
    border: 1.5px solid transparent !important;
    background:
        linear-gradient(var(--surface), var(--surface)) padding-box,
        conic-gradient(from var(--a), var(--g1), var(--g2), var(--g3), var(--g4), var(--g1)) border-box !important;
    animation: spin 4s linear infinite;
    box-shadow: 0 0 0 4px rgba(124,77,255,0.10), 0 20px 50px rgba(124,77,255,0.28);
    transform: translateY(-2px);
}
[data-testid="stChatInputSubmitButton"] {
    background: linear-gradient(135deg, var(--g1), var(--g2)) !important; color: #fff !important;
    border-radius: 50% !important; transition: transform 0.2s ease, opacity 0.2s ease;
}
[data-testid="stChatInputSubmitButton"]:hover:not(:disabled) { transform: scale(1.12) rotate(-8deg); }
[data-testid="stChatInputSubmitButton"]:disabled { background: var(--card-hover) !important; color: var(--faint) !important; }

/* ---- FILE UPLOAD BUTTON ---- */
[data-testid="stChatInput"] [data-testid="stChatInputFileUploadButton"],
[data-testid="stChatInput"] [data-testid="stChatInputFileUploadButton"] *,
[data-testid="stChatInput"] button[aria-label*="file" i],
[data-testid="stChatInput"] button[aria-label*="file" i] *,
[data-testid="stChatInput"] button[title*="file" i],
[data-testid="stChatInput"] button[title*="file" i] * {
    color: var(--muted) !important;
    -webkit-text-fill-color: var(--muted) !important;
    opacity: 1 !important;
    visibility: visible !important;
}
[data-testid="stChatInput"] [data-testid="stChatInputFileUploadButton"]:hover,
[data-testid="stChatInput"] [data-testid="stChatInputFileUploadButton"]:hover *,
[data-testid="stChatInput"] button[aria-label*="file" i]:hover,
[data-testid="stChatInput"] button[aria-label*="file" i]:hover * {
    color: var(--accent) !important;
    -webkit-text-fill-color: var(--accent) !important;
    background: var(--card-hover) !important;
    border-radius: 50%;
}

[data-testid="stChatInput"] [data-testid="stChatInputFileUploadButton"] + div,
[data-testid="stChatInput"] [data-baseweb="tag"] {
    background: var(--card) !important;
    border: 1px solid var(--line) !important;
    color: var(--text) !important;
    border-radius: 999px !important;
}
[data-testid="stChatInput"] [data-baseweb="tag"] * { color: var(--text) !important; }

/* ================= login ================= */
.login-wrap { display: flex; flex-direction: column; align-items: center; text-align: center; gap: 0.9rem; margin: 6vh 0 1.5rem 0; }
.login-wrap .orb { --s: 88px; animation: orbIn 0.9s cubic-bezier(0.2, 0.9, 0.3, 1.2) both, spin 7s linear infinite, breathe 4.5s ease-in-out 1s infinite; }
.login-title {
    font-family: 'Bricolage Grotesque', 'DM Sans', sans-serif; font-size: 2.3rem; font-weight: 700;
    letter-spacing: -0.035em; color: var(--text); animation: riseIn 0.6s 0.1s both;
}
.login-sub { color: var(--muted); animation: riseIn 0.6s 0.2s both; }
[data-testid="stForm"] {
    background: var(--panel); backdrop-filter: blur(16px); border: 1px solid var(--line);
    border-radius: 22px; box-shadow: var(--shadow); padding: 1.6rem; animation: riseIn 0.7s 0.25s both;
}
.stFormSubmitButton button {
    width: 100%; background: linear-gradient(120deg, var(--g1), var(--g2)) !important; border: none !important;
    border-radius: 14px !important; padding: 0.62rem !important; box-shadow: 0 10px 24px rgba(124,77,255,0.35);
    transition: transform 0.2s ease, box-shadow 0.2s ease;
}
.stFormSubmitButton button, .stFormSubmitButton button * { color: #fff !important; font-weight: 600; }
.stFormSubmitButton button:hover { transform: translateY(-2px); box-shadow: 0 14px 30px rgba(124,77,255,0.5); }

/* ================= attached files ================= */
.file-chips { display: flex; flex-wrap: wrap; gap: 0.35rem; margin: 0 0 0.45rem 0; }
.file-chip {
    display: inline-flex; align-items: center; gap: 0.35rem; max-width: 240px;
    padding: 0.2rem 0.7rem; border-radius: 999px; font-size: 0.8rem; font-weight: 500;
    background: rgba(255,255,255,0.22); overflow: hidden; text-overflow: ellipsis; white-space: nowrap;
}
[data-testid="stChatMessage"] [data-testid="stImage"] img { border-radius: 14px; margin-bottom: 0.4rem; }

/* ================= motion ================= */
@keyframes spin { to { --a: 360deg; } }
@keyframes breathe { 0%, 100% { transform: scale(1); } 50% { transform: scale(1.07); } }
@keyframes drift { from { transform: translate3d(0, 0, 0) scale(1); } to { transform: translate3d(4%, 3%, 0) scale(1.12) rotate(6deg); } }
@keyframes riseIn { from { opacity: 0; transform: translateY(14px); } to { opacity: 1; transform: none; } }
@keyframes bubbleIn { from { opacity: 0; transform: translateY(10px) scale(0.92); } to { opacity: 1; transform: none; } }
@keyframes orbIn { from { opacity: 0; transform: scale(0.3) rotate(-90deg); } to { opacity: 1; transform: none; } }
@keyframes sheen { from { background-position: 100% 0; } to { background-position: 0% 0; } }
@keyframes shine { to { left: 130%; } }
@keyframes wave { 0%, 80%, 100% { transform: translateY(0); opacity: 0.45; } 40% { transform: translateY(-6px); opacity: 1; } }
@keyframes shimmer { from { background-position: 100% 0; } to { background-position: -100% 0; } }

:focus-visible { outline: 2px solid var(--accent); outline-offset: 2px; }

@media (max-width: 640px) {
    .hero { margin-top: 3vh; }
    [data-testid="stChatMessage"]:has([data-testid="stChatMessageAvatarUser"]) { max-width: 92%; }
}
@media (prefers-reduced-motion: reduce) {
    *, *::before, *::after { animation: none !important; transition: none !important; }
}

/* ================= INPUT & SEARCH BAR FIX ================= */
/* Search bar aur text inputs ka text hamesha visible rahe */

/* Har input ka base background + text color */
[data-baseweb="input"],
[data-testid="stTextInput"] [data-baseweb="input"],
[data-testid="stSearchInput"] [data-baseweb="input"] {
    background-color: var(--card) !important;
    border: 1px solid var(--line) !important;
    border-radius: 12px !important;
}

/* Input ke andar ka actual <input> element */
[data-baseweb="input"] input,
[data-testid="stTextInput"] input,
[data-testid="stSearchInput"] input,
[data-baseweb="base-input"] input {
    color: var(--text) !important;
    -webkit-text-fill-color: var(--text) !important;
    background: transparent !important;
    caret-color: var(--accent) !important;
    opacity: 1 !important;
}

/* Placeholder text */
[data-baseweb="input"] input::placeholder,
[data-testid="stTextInput"] input::placeholder,
[data-testid="stSearchInput"] input::placeholder {
    color: var(--faint) !important;
    -webkit-text-fill-color: var(--faint) !important;
    opacity: 0.9 !important;
}

/* Focus state */
[data-baseweb="input"]:focus-within,
[data-testid="stTextInput"] [data-baseweb="input"]:focus-within,
[data-testid="stSearchInput"] [data-baseweb="input"]:focus-within {
    border-color: var(--accent) !important;
    box-shadow: 0 0 0 3px rgba(124,77,255,0.2) !important;
    background-color: var(--surface) !important;
}

/* Popover ke andar input ka white background fix */
[data-testid="stPopoverBody"] [data-baseweb="input"] {
    background-color: var(--surface) !important;
}
[data-testid="stPopoverBody"] input {
    color: var(--text) !important;
    -webkit-text-fill-color: var(--text) !important;
}

/* ================= THEME OVERRIDES (beat config.toml dark base) ================= */
.stApp [data-testid="stSidebar"] [data-baseweb="input"],
.stApp [data-testid="stSidebar"] [data-baseweb="base-input"],
.stApp [data-testid="stSidebar"] [data-baseweb="textarea"],
.stApp [data-testid="stSidebar"] [data-baseweb="textarea"] > div,
.stApp [data-testid="stSidebar"] [data-baseweb="select"] > div {
    background-color: var(--card) !important; background: var(--card) !important;
    border-color: var(--line) !important; border-radius: 12px !important;
}
.stApp [data-testid="stSidebar"] input,
.stApp [data-testid="stSidebar"] textarea {
    background: transparent !important; color: var(--text) !important;
    -webkit-text-fill-color: var(--text) !important; caret-color: var(--accent) !important; opacity: 1 !important;
}
.stApp [data-testid="stSidebar"] input::placeholder,
.stApp [data-testid="stSidebar"] textarea::placeholder {
    color: var(--faint) !important; -webkit-text-fill-color: var(--faint) !important; opacity: 1 !important;
}
.stApp [data-testid="stExpander"] details,
.stApp [data-testid="stExpander"] summary,
.stApp [data-testid="stExpander"] [data-testid="stExpanderDetails"] {
    background: transparent !important; color: var(--text) !important;
}
.stApp [data-testid="stExpander"] summary:hover { background: var(--card-hover) !important; }
.stApp [data-testid="stExpander"] summary *, .stApp [data-testid="stExpander"] summary svg {
    color: var(--text) !important; fill: var(--text) !important; -webkit-text-fill-color: var(--text) !important;
}
.stApp [data-testid="stToggle"] label p,
.stApp [data-testid="stCheckbox"] label p {
    color: var(--text) !important;
}

/* ---- input wrappers (Streamlit paints the dark fill on these) ---- */
.stApp [data-testid="stSidebar"] [data-testid="stTextInputRootElement"],
.stApp [data-testid="stSidebar"] [data-testid="stTextInputRootElement"] > div,
.stApp [data-testid="stSidebar"] [data-testid="stTextAreaRootElement"],
.stApp [data-testid="stSidebar"] [data-testid="stTextAreaRootElement"] > div,
.stApp [data-testid="stSidebar"] [data-testid="stTextInput"] > div > div,
.stApp [data-testid="stSidebar"] [data-testid="stTextArea"] > div > div,
.stApp [data-testid="stSidebar"] [data-testid="stTextArea"] [data-baseweb] ,
.stApp [data-testid="stSidebar"] [data-testid="stTextInput"] [data-baseweb] {
    background: var(--card) !important; background-color: var(--card) !important; background-image: none !important;
    border-color: var(--line) !important; border-radius: 12px !important;
}
.stApp [data-testid="stSidebar"] [data-testid="stTextInput"] input,
.stApp [data-testid="stSidebar"] [data-testid="stTextArea"] textarea {
    background: transparent !important; background-color: transparent !important;
}

/* ================= scrollable sidebar ================= */
[data-testid="stSidebar"] { overflow: hidden !important; }
[data-testid="stSidebarContent"] {
    height: 100% !important; max-height: 100vh !important;
    overflow-y: auto !important; overflow-x: hidden !important;
    scrollbar-width: thin; scrollbar-color: var(--line) transparent;
}
[data-testid="stSidebarUserContent"] { padding-bottom: 2rem !important; overflow: visible !important; }
[data-testid="stSidebarContent"]::-webkit-scrollbar { width: 6px; }
[data-testid="stSidebarContent"]::-webkit-scrollbar-thumb { background: var(--line); border-radius: 6px; }
"""


def inject_css(dark: bool):
    t = THEMES["dark" if dark else "light"]
    variables = "".join(f"--{k}:{v};" for k, v in t.items() if k != "scheme")
    css = f":root{{{variables}color-scheme:{t['scheme']};}}\n" + BASE_CSS
    st.markdown(f"<style>{css}</style>", unsafe_allow_html=True)


# ---------- Session state ----------
st.session_state.setdefault("chats", {})
st.session_state.setdefault("active", None)
st.session_state.setdefault("api_key", "")
st.session_state.setdefault("user_name", "")
st.session_state.setdefault("theme_dark", True)
st.session_state.setdefault("popover_cid", None)


def _set_theme():
    st.session_state.theme_dark = st.session_state.dark_toggle


inject_css(st.session_state.theme_dark)


# ---------- Storage ----------
def db_path():
    digest = hashlib.sha256(st.session_state.api_key.encode()).hexdigest()[:16]
    return f"chats_{digest}.json"


def load_chats():
    path = db_path()
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}


def save_chats():
    with open(db_path(), "w", encoding="utf-8") as f:
        json.dump(st.session_state.chats, f, ensure_ascii=False, indent=2)


def new_chat():
    cid = str(uuid.uuid4())
    st.session_state.chats[cid] = {
        "title": "New chat",
        "system": DEFAULT_SYSTEM,
        "messages": [],
        "updated": datetime.now().isoformat(),
    }
    st.session_state.active = cid
    save_chats()
    return cid


def sorted_chat_ids():
    chats = st.session_state.chats
    return sorted(chats, key=lambda c: chats[c].get("updated", ""), reverse=True)


def group_label(iso):
    try:
        days = (datetime.now().date() - datetime.fromisoformat(iso).date()).days
    except Exception:
        return "Older"
    if days <= 0:
        return "Today"
    if days == 1:
        return "Yesterday"
    if days <= 7:
        return "Previous 7 days"
    if days <= 30:
        return "Previous 30 days"
    return "Older"


def short(text, n=30):
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


# ---------- File attachments ----------
IMAGE_TYPES = ["png", "jpg", "jpeg", "webp", "gif"]
DOC_TYPES = ["pdf", "docx", "xlsx"]
TEXT_TYPES = [
    "txt", "md", "csv", "tsv", "json", "xml", "yaml", "yml", "toml", "ini", "log",
    "py", "js", "ts", "jsx", "tsx", "html", "css", "java", "c", "cpp", "h", "cs",
    "go", "rs", "php", "rb", "sql", "sh",
]
ACCEPTED_TYPES = IMAGE_TYPES + DOC_TYPES + TEXT_TYPES
MAX_FILE_MB = 10
MAX_CHARS_PER_FILE = 60_000
MAX_CHARS_TOTAL = 150_000
DEFAULT_FILE_PROMPT = "Please summarize the attached file(s)."


def _extract_text(ext, data):
    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = [f"[Page {i}]\n{(p.extract_text() or '').strip()}" for i, p in enumerate(reader.pages, 1)]
        return "\n\n".join(pages)
    if ext == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [para.text for para in doc.paragraphs if para.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    parts.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(parts)
    return data.decode("utf-8", errors="replace")


def read_uploaded(f):
    name = f.name
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    data = f.getvalue()
    if len(data) > MAX_FILE_MB * 1024 * 1024:
        return {"name": name, "kind": "error", "text": f"larger than {MAX_FILE_MB} MB"}
    if ext in IMAGE_TYPES:
        mime = "image/jpeg" if ext in ("jpg", "jpeg") else f"image/{ext}"
        return {"name": name, "kind": "image", "data": f"data:{mime};base64," + base64.b64encode(data).decode()}
    try:
        text = _extract_text(ext, data).strip()
    except Exception as e:
        return {"name": name, "kind": "error", "text": f"couldn't read it ({e})"}
    if not text:
        return {"name": name, "kind": "error", "text": "no readable text found (is it a scanned file?)"}
    if len(text) > MAX_CHARS_PER_FILE:
        text = text[:MAX_CHARS_PER_FILE] + "\n[... file truncated because it is too long ...]"
    return {"name": name, "kind": "text", "text": text}


def unpack_input(value):
    if value is None:
        return "", []
    if isinstance(value, str):
        return value, []
    text = getattr(value, "text", None)
    files = getattr(value, "files", None)
    if text is None:
        text = value.get("text", "")
    if files is None:
        files = value.get("files", [])
    return text or "", list(files or [])


def build_user_content(m, budget):
    files = m.get("files") or []
    if not files:
        return m["content"], budget
    blocks, images = [], []
    for f in files:
        if f["kind"] == "image":
            images.append(f)
        elif f.get("text"):
            if len(f["text"]) <= budget:
                budget -= len(f["text"])
                blocks.append(f"--- File: {f['name']} ---\n{f['text']}\n--- End of {f['name']} ---")
            else:
                blocks.append(f"--- File: {f['name']} --- (content left out: too much file text in this chat)")
    text = m["content"]
    if blocks:
        text = (
            "The user attached these file(s). Use them to answer.\n\n"
            + "\n\n".join(blocks)
            + f"\n\nUser message: {m['content']}"
        )
    if images:
        content = [{"type": "text", "text": text}]
        content += [{"type": "image_url", "image_url": {"url": f["data"]}} for f in images]
        return content, budget
    return text, budget


def build_history(chat):
    budget = MAX_CHARS_TOTAL
    out = []
    for m in reversed(chat["messages"]):
        if m["role"] == "user":
            content, budget = build_user_content(m, budget)
            out.append(HumanMessage(content=content))
        else:
            out.append(AIMessage(content=m["content"]))
    out.append(SystemMessage(content=chat["system"] or DEFAULT_SYSTEM))
    return out[::-1]


def render_user_message(m):
    files = m.get("files") or []
    chips = "".join(
        f"<span class='file-chip'>📎 {html.escape(f['name'])}</span>" for f in files
    )
    if chips:
        st.markdown(f"<div class='file-chips'>{chips}</div>", unsafe_allow_html=True)
    for f in files:
        if f["kind"] == "image":
            st.image(base64.b64decode(f["data"].split(",", 1)[1]), width=240)
    st.markdown(m["content"])


# ---------- Login gate ----------
if not st.session_state.api_key:
    st.markdown(
        "<div class='login-wrap'><div class='orb'></div>"
        f"<div class='login-title'>{APP_NAME}</div>"
        "<div class='login-sub'>Enter your OpenAI API key to start chatting</div></div>",
        unsafe_allow_html=True,
    )
    with st.form("login_form"):
        name_in = st.text_input("Your name (optional)", placeholder="e.g. Fatima")
        key_in = st.text_input("OpenAI API key", type="password", placeholder="sk-...")
        enter = st.form_submit_button("Start chatting")
    if enter:
        key_in = key_in.strip()
        if not key_in:
            st.error("Enter your API key to continue.")
        else:
            try:
                from openai import OpenAI
                with st.spinner("Checking your key..."):
                    OpenAI(api_key=key_in).models.list()
                st.session_state.api_key = key_in
                st.session_state.user_name = name_in.strip()
                st.session_state.chats = load_chats()
                ids = sorted_chat_ids()
                st.session_state.active = ids[0] if ids else None
                st.rerun()
            except Exception:
                st.error("That API key didn't work. Check it for typos and try again.")
    st.caption("Your key stays in this session. It is never saved to a file.")
    st.stop()

# ---------- Sidebar ----------
with st.sidebar:
    st.markdown(
        f"<div class='sb-brand'><div class='orb sm'></div><div class='sb-name'>{APP_NAME}</div></div>",
        unsafe_allow_html=True,
    )

    if st.button("New chat", key="new_chat_btn", icon=":material/add:"):
        new_chat()
        st.session_state.popover_cid = None
        st.rerun()

    search = st.text_input(
        "Search chats", placeholder="Search chats", label_visibility="collapsed", key="chat_search"
    ).strip().lower()

    # ----- chat history, grouped by date -----
    with st.container(height=260, border=False):
        shown = 0
        last_group = None
        for cid in sorted_chat_ids():
            chat_item = st.session_state.chats[cid]
            if search and search not in chat_item["title"].lower():
                continue
            shown += 1
            group = group_label(chat_item.get("updated", ""))
            if group != last_group:
                st.markdown(f"<div class='sb-group'>{group}</div>", unsafe_allow_html=True)
                last_group = group

            is_active = cid == st.session_state.active
            c1, c2 = st.columns([7, 1.6])
            with c1:
                key = "active_chat_btn" if is_active else f"sel_{cid}"
                if st.button(short(chat_item["title"]), key=key):
                    st.session_state.active = cid
                    st.session_state.popover_cid = None
                    st.rerun()
            with c2:
                with st.popover("⋯"):
                    title_hash = hashlib.md5(chat_item["title"].encode()).hexdigest()[:8]
                    new_title = st.text_input(
                        "Chat name",
                        value=chat_item["title"],
                        key=f"rn_{cid}_{title_hash}",
                        label_visibility="collapsed",
                    )
                    if st.button("Save name", key=f"save_{cid}", icon=":material/check:"):
                        chat_item["title"] = new_title.strip() or chat_item["title"]
                        save_chats()
                        st.rerun()
                    if st.button("Delete chat", key=f"del_{cid}", icon=":material/delete:"):
                        del st.session_state.chats[cid]
                        if st.session_state.active == cid:
                            ids = sorted_chat_ids()
                            st.session_state.active = ids[0] if ids else None
                        save_chats()
                        st.rerun()

        if shown == 0:
            if search:
                st.caption(f"No chats match “{search}”.")
            else:
                st.caption("No chats yet. Send a message to start one.")

    with st.expander("Model & appearance"):
        model_name = st.text_input("Model", value=DEFAULT_MODEL, key="model_name")
        st.toggle(
            "Dark mode",
            value=st.session_state.theme_dark,
            key="dark_toggle",
            on_change=_set_theme,
        )
        active_id = st.session_state.active
        if active_id and active_id in st.session_state.chats:
            sys_val = st.text_area(
                "System role",
                value=st.session_state.chats[active_id]["system"],
                key=f"sys_{active_id}",
                height=110,
            )
            if sys_val != st.session_state.chats[active_id]["system"]:
                st.session_state.chats[active_id]["system"] = sys_val
                save_chats()
        else:
            st.caption("Open a chat to set its system role.")

    who = st.session_state.user_name or "Guest"
    st.markdown(
        f"<div class='user-chip'><div class='av'>{html.escape(who[:1].upper())}</div>"
        f"<div><div class='nm'>{html.escape(who)}</div><div class='sub'>Signed in with API key</div></div></div>",
        unsafe_allow_html=True,
    )
    if st.button("Log out", key="logout_btn", icon=":material/logout:"):
        st.session_state.api_key = ""
        st.session_state.user_name = ""
        st.session_state.chats = {}
        st.session_state.active = None
        st.session_state.popover_cid = None
        st.rerun()

# ---------- Main area ----------
active_id = st.session_state.active
chat = st.session_state.chats.get(active_id) if active_id else None
has_messages = bool(chat and chat["messages"])

pending = st.session_state.pop("pending_prompt", None)
regen = bool(st.session_state.pop("regen_flag", False)) and has_messages and chat["messages"][-1]["role"] == "user"

typed = None
holder = st.empty()

if has_messages:
    typed = st.chat_input(
        f"Message {APP_NAME}...", key="input_bottom",
        accept_file="multiple", file_type=ACCEPTED_TYPES,
    )
    typed_text, typed_files = unpack_input(typed)
    if typed_files and not typed_text.strip():
        typed_text = DEFAULT_FILE_PROMPT
    prompt = typed_text or pending
    prompt_files = typed_files if typed_text else []
    with holder.container():
        for m in chat["messages"]:
            avatar = USER_AVATAR if m["role"] == "user" else BOT_AVATAR
            with st.chat_message(m["role"], avatar=avatar):
                if m["role"] == "user":
                    render_user_message(m)
                else:
                    st.markdown(m["content"])
        if not prompt and not regen and chat["messages"][-1]["role"] == "assistant":
            if st.button("Regenerate", key="regen_btn", icon=":material/refresh:"):
                chat["messages"].pop()
                st.session_state.regen_flag = True
                save_chats()
                st.rerun()
else:
    hour = datetime.now().hour
    part = "Good morning" if 5 <= hour < 12 else "Good afternoon" if 12 <= hour < 18 else "Good evening"
    name = html.escape(st.session_state.user_name)
    title = f"{part}, {name}" if name else part
    with holder.container():
        st.markdown(
            f"<div class='hero'><div class='orb'></div>"
            f"<div class='hero-title'>{title}</div>"
            "<p>What are we working on today?</p></div>",
            unsafe_allow_html=True,
        )
        typed = st.chat_input(
            f"Message {APP_NAME}...", key="input_center",
            accept_file="multiple", file_type=ACCEPTED_TYPES,
        )

        cats = list(SUGGESTIONS)
        cat = st.pills(
            "Topics", cats, default="Popular", key="sg_cat", label_visibility="collapsed"
        ) or "Popular"
        items = SUGGESTIONS[cat]
        for row in range(0, len(items), 2):
            cols = st.columns(2)
            for j, col in enumerate(cols):
                i = row + j
                if i >= len(items):
                    continue
                icon, label, full_prompt = items[i]
                with col:
                    if st.button(label, key=f"sg_card_{i}_{cat.lower()}", icon=f":material/{icon}:"):
                        st.session_state.pending_prompt = full_prompt
                        st.rerun()
        st.markdown(
            f"<div class='fine'>{APP_NAME} can make mistakes. Check important info.</div>",
            unsafe_allow_html=True,
        )
    typed_text, typed_files = unpack_input(typed)
    if typed_files and not typed_text.strip():
        typed_text = DEFAULT_FILE_PROMPT
    prompt = typed_text or pending
    prompt_files = typed_files if typed_text else []

if prompt or regen:
    if not chat:
        active_id = new_chat()
        chat = st.session_state.chats[active_id]

    if not has_messages and prompt:
        holder.empty()

    added_user_message = False
    if prompt:
        attachments = []
        if prompt_files:
            with st.spinner("Reading your files..."):
                for f in prompt_files:
                    result = read_uploaded(f)
                    if result["kind"] == "error":
                        st.toast(f"Skipped {result['name']}: {result['text']}", icon=":material/error:")
                    else:
                        attachments.append(result)
        user_msg = {"role": "user", "content": prompt}
        if attachments:
            user_msg["files"] = attachments
        chat["messages"].append(user_msg)
        added_user_message = True
        if chat["title"] == "New chat":
            if attachments and prompt == DEFAULT_FILE_PROMPT:
                chat["title"] = attachments[0]["name"][:40]
            else:
                chat["title"] = prompt.strip().replace("\n", " ")[:40]
        with st.container(key="live_user"):
            with st.chat_message("user", avatar=USER_AVATAR):
                render_user_message(user_msg)

    history = build_history(chat)

    try:
        llm = ChatOpenAI(model=model_name, api_key=st.session_state.api_key)
        with st.container(key="live_bot"):
            with st.chat_message("assistant", avatar=BOT_AVATAR):
                typing = st.empty()
                typing.markdown(
                    "<div class='thinking'><div class='dots'><i></i><i></i><i></i></div>"
                    "<span class='label'>Thinking</span></div>",
                    unsafe_allow_html=True,
                )

                def token_stream():
                    first = True
                    for c in llm.stream(history):
                        if c.content:
                            if first:
                                typing.empty()
                                first = False
                            yield c.content

                reply = st.write_stream(token_stream())
                typing.empty()
        chat["messages"].append({"role": "assistant", "content": reply})
        chat["updated"] = datetime.now().isoformat()
        save_chats()
        st.rerun()
    except Exception as e:
        if added_user_message:
            chat["messages"].pop()
            if not chat["messages"]:
                chat["title"] = "New chat"
        save_chats()
        st.error(f"Couldn't get a reply: {e}")